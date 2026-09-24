# GET /api/reports/daily-report/excel/?date=2026-09-24

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter  # Safer way to get column letters

def generate_daily_report_excel(report_data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Report"

    headers = [
        "School Code", "School Name",
        "Bun Demand", "Bun Delivered", "Bun Shortfall",
        "Egg Demand", "Egg Delivered", "Egg Shortfall",
        "Banana Demand", "Banana Delivered", "Banana Shortfall",
    ]
    ws.append(headers)

    # Style headers bold
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # 1. Append all school rows first
    for row in report_data["schools"]:
        ws.append([
            row["school_code"],
            row["school_name"],
            row["bun_demand"],
            row["bun_delivered"],
            row["bun_shortfall"],
            row["egg_demand"],
            row["egg_delivered"],
            row["egg_shortfall"],
            row["banana_demand"],
            row["banana_delivered"],
            row["banana_shortfall"],
        ])

    # 2. FIXED: Move totals OUTSIDE the loop so they appear once at the very end
    totals = report_data["totals"]
    ws.append([])  # Blank spacer row
    
    ws.append([
        "",
        "UPAZILA TOTAL",
        totals["bun_demand"],
        totals["bun_delivered"],
        totals["bun_shortfall"],
        totals["egg_demand"],
        totals["egg_delivered"],
        totals["egg_shortfall"],
        totals["banana_demand"],
        totals["banana_delivered"],
        totals["banana_shortfall"],
    ])

    # Style the Upazila Total row bold (it's the last row in the worksheet)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    # 3. Dynamic Column Auto-Width Fit
    for col in ws.columns:
        max_length = 0
        # Use openpyxl utility to safely fetch column letter
        column_letter = get_column_letter(col[0].column)
        
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
                
        ws.column_dimensions[column_letter].width = max(max_length + 3, 12)

    # Save to memory stream
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output



# from io import BytesIO

# from openpyxl import Workbook
# from openpyxl.styles import Font


# def generate_daily_report_excel(report_data):

#     wb = Workbook()

#     ws = wb.active

#     ws.title = "Daily Report"

#     headers = [

#         "School Code",
#         "School Name",

#         "Bun Demand",
#         "Bun Delivered",
#         "Bun Shortfall",

#         "Egg Demand",
#         "Egg Delivered",
#         "Egg Shortfall",

#         "Banana Demand",
#         "Banana Delivered",
#         "Banana Shortfall",
#     ]

#     ws.append(headers)

#     for cell in ws[1]:
#         cell.font = Font(bold=True)

#     for row in report_data["schools"]:

#         ws.append([

#             row["school_code"],
#             row["school_name"],

#             row["bun_demand"],
#             row["bun_delivered"],
#             row["bun_shortfall"],

#             row["egg_demand"],
#             row["egg_delivered"],
#             row["egg_shortfall"],

#             row["banana_demand"],
#             row["banana_delivered"],
#             row["banana_shortfall"],
#         ])

#         totals = report_data["totals"]

#         ws.append([])

#         ws.append([

#             "",
#             "UPAZILA TOTAL",

#             totals["bun_demand"],
#             totals["bun_delivered"],
#             totals["bun_shortfall"],

#             totals["egg_demand"],
#             totals["egg_delivered"],
#             totals["egg_shortfall"],

#             totals["banana_demand"],
#             totals["banana_delivered"],
#             totals["banana_shortfall"],
#         ])

#     for column in ws.columns:

#         max_length = 0

#         column_letter = column[0].column_letter

#         for cell in column:

#             try:
#                 max_length = max(
#                     max_length,
#                     len(str(cell.value))
#                 )
#             except:
#                 pass

#         ws.column_dimensions[
#             column_letter
#         ].width = max_length + 5

#     output = BytesIO()

#     wb.save(output)

#     output.seek(0)

#     return output